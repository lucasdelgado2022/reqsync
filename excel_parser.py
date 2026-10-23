from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


ROOT_SPEC_ID = "Estructura de Requerimientos automaticos"
REQUIRED_COLUMNS = {"internalid", "content", "parent_ids", "type"}


@dataclass(frozen=True)
class RequirementNode:
    internalid: str
    object_type: str
    content: str


def parse_requirement_structure(
    workbook_path: str | Path, sheet_name: str | None = None
) -> tuple[dict[str, RequirementNode], set[tuple[str, str]]]:
    path = Path(workbook_path)
    if not path.is_file():
        raise FileNotFoundError(f"Workbook does not exist: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name is None:
            sheet = workbook.active
        elif sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Worksheet {sheet_name!r} not found; available sheets: "
                f"{', '.join(workbook.sheetnames)}"
            )
        else:
            sheet = workbook[sheet_name]

        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration as error:
            raise ValueError("Workbook worksheet is empty") from error

        headers = [
            str(value).strip().lower() if value is not None else ""
            for value in header_row
        ]
        non_empty_headers = [header for header in headers if header]
        if len(non_empty_headers) != len(set(non_empty_headers)):
            raise ValueError("Workbook contains duplicate column headers")

        missing = REQUIRED_COLUMNS.difference(headers)
        if missing:
            raise ValueError(
                "Workbook is missing required columns: "
                + ", ".join(sorted(missing))
            )

        index = {name: headers.index(name) for name in REQUIRED_COLUMNS}
        nodes: dict[str, RequirementNode] = {}
        parent_ids: dict[str, tuple[str, ...]] = {}

        for row_number, row in enumerate(rows, start=2):
            values = tuple(row)
            if not any(value is not None and str(value).strip() for value in values):
                continue

            raw_id = (
                values[index["internalid"]]
                if index["internalid"] < len(values)
                else None
            )
            if raw_id is None or not str(raw_id).strip():
                raise ValueError(f"Row {row_number}: internalid is required")
            internalid = str(raw_id).strip()
            if internalid == ROOT_SPEC_ID:
                raise ValueError(
                    f"Row {row_number}: {ROOT_SPEC_ID!r} is reserved for the "
                    "root specification"
                )
            if internalid in nodes:
                raise ValueError(
                    f"Row {row_number}: duplicate internalid {internalid!r}"
                )

            raw_type = values[index["type"]] if index["type"] < len(values) else None
            type_code = str(raw_type).strip().upper() if raw_type is not None else ""
            if type_code not in {"", "C", "R"}:
                raise ValueError(
                    f"Row {row_number}: type must be C, R, or blank; got {type_code!r}"
                )
            object_type = "Chapter" if type_code == "C" else "Requirement"

            raw_content = (
                values[index["content"]] if index["content"] < len(values) else None
            )
            content = "" if raw_content is None else str(raw_content)
            raw_parents = (
                values[index["parent_ids"]]
                if index["parent_ids"] < len(values)
                else None
            )
            parents = tuple(
                dict.fromkeys(
                    parent.strip()
                    for parent in str(raw_parents or "").split(";")
                    if parent.strip()
                )
            )

            nodes[internalid] = RequirementNode(internalid, object_type, content)
            parent_ids[internalid] = parents

        if not nodes:
            raise ValueError("Workbook contains no requirement or chapter rows")

        edges: set[tuple[str, str]] = set()
        for child_id, parents in parent_ids.items():
            child = nodes[child_id]
            if not parents:
                if child.object_type != "Chapter":
                    raise ValueError(
                        f"Top-level requirement {child_id!r} is not allowed; "
                        "the specification can contain chapters only"
                    )
                edges.add((ROOT_SPEC_ID, child_id))
                continue

            for parent_id in parents:
                if parent_id not in nodes:
                    raise ValueError(
                        f"Row for {child_id!r}: unknown parent {parent_id!r}"
                    )
                parent = nodes[parent_id]
                if child.object_type == "Chapter" and parent.object_type != "Chapter":
                    raise ValueError(
                        f"Chapter {child_id!r} cannot be contained by "
                        f"{parent.object_type.lower()} {parent_id!r}"
                    )
                edges.add((parent_id, child_id))

        _validate_acyclic(nodes, edges)
        return nodes, edges
    finally:
        workbook.close()


def _validate_acyclic(
    nodes: dict[str, RequirementNode], edges: set[tuple[str, str]]
) -> None:
    children: dict[str, list[str]] = {internalid: [] for internalid in nodes}
    for parent_id, child_id in edges:
        if parent_id in children:
            children[parent_id].append(child_id)

    states: dict[str, int] = {}

    def visit(internalid: str) -> None:
        state = states.get(internalid, 0)
        if state == 1:
            raise ValueError(f"Hierarchy contains a cycle involving {internalid!r}")
        if state == 2:
            return
        states[internalid] = 1
        for child_id in children[internalid]:
            visit(child_id)
        states[internalid] = 2

    for internalid in nodes:
        visit(internalid)
